#!/usr/bin/env python
import argparse
import datetime as dt
import io
import json
import re
import sqlite3
import sys
import time
import zipfile
from pathlib import Path

import openpyxl

from normalize_sqlite import SCHEMA


DEFAULT_ZIP = Path(__file__).resolve().parent / ".." / "2025-07-24 13-40-58.zip"
DEFAULT_DB = Path(__file__).resolve().parent / "vuzonline_2025_norm.sqlite"


HEADER_MAP = {
    "Порядковый номер": "ordinal",
    "ID участника": "applicant_id",
    "Приоритет конкурса": "competition_priority",
    "Подано согласие": "consent_text",
    "Сумма баллов": "score_total",
    "Баллы за ВИ": "exam_scores_text",
    "Баллы за ИД": "individual_score",
    "Статус": "status",
    "Дата выбора конкурсной группы по Москве": "selected_at_moscow",
}


def now_iso() -> str:
    return dt.datetime.now(dt.UTC).replace(microsecond=0).isoformat()


def as_text(value):
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def as_int(value):
    text = as_text(value)
    if text is None or text == "-":
        return None
    try:
        return int(float(text.replace(",", ".")))
    except ValueError:
        return None


def as_float(value):
    text = as_text(value)
    if text is None or text == "-":
        return None
    try:
        return float(text.replace(",", "."))
    except ValueError:
        return None


def has_consent(value) -> int:
    text = as_text(value)
    return 0 if not text or text == "-" else 1


def parse_competition_path(source_path: str) -> dict:
    parts = source_path.split("/")
    file_name = parts[-1]
    stem = file_name.removesuffix(".xlsx")
    quota_type = stem.split(" [", 1)[0]
    seats = None
    match = re.search(r"\[(\d+)\]", stem)
    if match:
        seats = int(match.group(1))
    payload = {
        "source_path": source_path,
        "file_name": file_name,
        "program": None,
        "education_level": None,
        "city": None,
        "university": None,
        "study_form": None,
        "quota_type": quota_type,
        "seats": seats,
        "path_json": json.dumps(parts, ensure_ascii=False),
    }
    if len(parts) >= 8:
        payload.update(
            {
                "program": parts[2],
                "education_level": parts[3],
                "city": parts[4],
                "university": parts[5],
                "study_form": parts[6],
            }
        )
    return payload


def init_db(conn: sqlite3.Connection) -> None:
    conn.executescript(SCHEMA)
    conn.execute(
        """
        INSERT OR IGNORE INTO users(username, password_text, role, applicant_id, is_active, created_at, updated_at)
        VALUES (?, ?, 'admin', NULL, 1, ?, ?)
        """,
        ("prointdata@ya.ru", "admin", now_iso(), now_iso()),
    )
    conn.execute("UPDATE users SET is_active = 0, updated_at = ? WHERE username = 'admin'", (now_iso(),))
    conn.commit()


def get_or_create_id(conn: sqlite3.Connection, table: str, value: str | None) -> int | None:
    if not value:
        return None
    conn.execute(f"INSERT OR IGNORE INTO {table}(name) VALUES (?)", (value,))
    row = conn.execute(f"SELECT id FROM {table} WHERE name = ?", (value,)).fetchone()
    return row[0] if row else None


def import_workbook_norm(conn: sqlite3.Connection, archive: zipfile.ZipFile, source_path: str) -> int:
    if conn.execute("SELECT 1 FROM competitions_core WHERE source_path = ?", (source_path,)).fetchone():
        return 0

    meta = parse_competition_path(source_path)
    program_id = get_or_create_id(conn, "programs", meta["program"])
    university_id = get_or_create_id(conn, "universities", meta["university"])
    cur = conn.execute(
        """
        INSERT INTO competitions_core(
            source_path, file_name, program_id, education_level, city, university_id,
            study_form, quota_type, seats, path_json, row_count, imported_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0, ?)
        """,
        (
            meta["source_path"],
            meta["file_name"],
            program_id,
            meta["education_level"],
            meta["city"],
            university_id,
            meta["study_form"],
            meta["quota_type"],
            meta["seats"],
            meta["path_json"],
            now_iso(),
        ),
    )
    competition_id = cur.lastrowid

    wb = openpyxl.load_workbook(io.BytesIO(archive.read(source_path)), read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        header_row = []
    headers = [as_text(cell) or f"column_{idx + 1}" for idx, cell in enumerate(header_row)]
    mapped = [HEADER_MAP.get(header) for header in headers]

    pending = []
    row_count = 0
    for row in rows:
        values = {field: None for field in HEADER_MAP.values()}
        for idx, field in enumerate(mapped):
            if field and idx < len(row):
                values[field] = row[idx]
        applicant_external_id = as_text(values["applicant_id"])
        if not applicant_external_id:
            continue
        pending.append(
            (
                applicant_external_id,
                as_int(values["ordinal"]),
                as_int(values["competition_priority"]),
                as_text(values["consent_text"]),
                has_consent(values["consent_text"]),
                as_float(values["score_total"]),
                as_text(values["exam_scores_text"]),
                as_float(values["individual_score"]),
                as_text(values["status"]),
                as_text(values["selected_at_moscow"]),
            )
        )
        row_count += 1

        if len(pending) >= 1000:
            flush_applications(conn, competition_id, pending)
            pending.clear()
    if pending:
        flush_applications(conn, competition_id, pending)

    conn.execute("UPDATE competitions_core SET row_count = ? WHERE id = ?", (row_count, competition_id))
    wb.close()
    return row_count


def flush_applications(conn: sqlite3.Connection, competition_id: int, rows: list[tuple]) -> None:
    applicant_ids = [(row[0],) for row in rows]
    conn.executemany("INSERT OR IGNORE INTO applicants(applicant_id) VALUES (?)", applicant_ids)
    placeholders = ",".join("?" for _ in applicant_ids)
    id_map = {
        row[0]: row[1]
        for row in conn.execute(
            f"SELECT applicant_id, id FROM applicants WHERE applicant_id IN ({placeholders})",
            [row[0] for row in rows],
        )
    }
    conn.executemany(
        """
        INSERT INTO applications(
            competition_id, applicant_id, ordinal, competition_priority,
            consent_text, has_consent, score_total, exam_scores_text,
            individual_score, status, selected_at_moscow
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        [
            (
                competition_id,
                id_map[row[0]],
                row[1],
                row[2],
                row[3],
                row[4],
                row[5],
                row[6],
                row[7],
                row[8],
                row[9],
            )
            for row in rows
        ],
    )


def main() -> int:
    parser = argparse.ArgumentParser(description="Import XLSX zip directly into normalized SQLite.")
    parser.add_argument("--zip", type=Path, default=DEFAULT_ZIP)
    parser.add_argument("--db", type=Path, default=DEFAULT_DB)
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--commit-every", type=int, default=1000)
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()

    if not args.zip.exists():
        print(f"zip not found: {args.zip}", file=sys.stderr)
        return 2
    if args.db.exists() and args.force:
        args.db.unlink()
    if args.db.exists():
        print(f"db exists, pass --force to overwrite: {args.db}", file=sys.stderr)
        return 2

    conn = sqlite3.connect(args.db)
    conn.execute("PRAGMA journal_mode = WAL")
    conn.execute("PRAGMA synchronous = NORMAL")
    conn.execute("PRAGMA temp_store = MEMORY")
    init_db(conn)

    started = time.time()
    with zipfile.ZipFile(args.zip) as archive:
        names = [name for name in archive.namelist() if name.lower().endswith(".xlsx")]
        if args.limit:
            names = names[: args.limit]
        total_rows = 0
        for idx, name in enumerate(names, 1):
            total_rows += import_workbook_norm(conn, archive, name)
            if idx % args.commit_every == 0:
                conn.commit()
                print(f"imported_files={idx}/{len(names)} rows={total_rows}", flush=True)
        conn.commit()

    stats = {
        "applicants": conn.execute("SELECT count(*) FROM applicants").fetchone()[0],
        "universities": conn.execute("SELECT count(*) FROM universities").fetchone()[0],
        "programs": conn.execute("SELECT count(*) FROM programs").fetchone()[0],
        "competitions": conn.execute("SELECT count(*) FROM competitions_core").fetchone()[0],
        "applications": conn.execute("SELECT count(*) FROM applications").fetchone()[0],
    }
    print(f"stats {stats}")
    print(f"elapsed_sec {time.time() - started:.1f}")
    conn.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
